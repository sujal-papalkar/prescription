import docx

doc = docx.Document("data/uploads/DM_RR_DM RR.docx")
print("Total tables in DOCX:", len(doc.tables))
print("Total paragraphs in DOCX:", len(doc.paragraphs))

for i, table in enumerate(doc.tables):
    first_row = [c.text.strip().replace('\n', ' ') for c in table.rows[0].cells]
    first_few_cells = " | ".join(first_row[:3])
    print(f"\n--- TABLE INDEX {i} (Parser labeled as 'Page {i+1}') ---")
    print(f"Header: {first_few_cells[:120]}")
    # print sample row
    if len(table.rows) > 1:
        sample_row = [c.text.strip().replace('\n', ' ') for c in table.rows[1].cells]
        print(f"Sample row: {' | '.join(sample_row[:3])[:120]}")
