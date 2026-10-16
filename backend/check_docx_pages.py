import docx

doc = docx.Document("data/uploads/DM_RR_DM RR.docx")

print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")

current_page = 1
for elem in doc.element.body:
    tag = elem.tag.split("}")[-1]
    if tag == "p":
        txt = "".join(elem.itertext()).strip()
        # Find page breaks
        breaks = elem.xpath('.//*[local-name()="lastRenderedPageBreak" or (local-name()="br" and @*[local-name()="type"]="page")]')
        if breaks:
            current_page += len(breaks)
            print(f"--- PAGE BREAK DETECTED -> NOW PAGE {current_page} ---")
        if txt and len(txt) > 20:
            print(f"[Page {current_page}] Paragraph: {txt[:80]}...")
    elif tag == "tbl":
        t = docx.table.Table(elem, doc)
        headers = [c.text.strip().replace("\n", " ") for c in t.rows[0].cells[:2]]
        print(f"*** TABLE ENCOUNTERED AT PAGE {current_page} *** Header: {headers}")
        # check page breaks inside table rows
        tbl_breaks = elem.xpath('.//*[local-name()="lastRenderedPageBreak" or (local-name()="br" and @*[local-name()="type"]="page")]')
        if tbl_breaks:
            print(f"    (Table contains {len(tbl_breaks)} page breaks, spans to Page {current_page + len(tbl_breaks)})")
            current_page += len(tbl_breaks)
